"""
build_subsystem_retrofit.py
===========================

Retrofits six-way CAV subsystem labels onto the completed tranche 1
annotation sheets.

Why this exists
---------------
R5 of docs/analysis_preregistration.md pre-commits to subsystem labels on
the 200-item reliability subset (100 per tranche). The calibration sheets
carried a `subsystem` column; the tranche 1 production sheets did not.
This script reads the three completed tranche 1 sheets and emits a second,
narrower pass that collects only the missing subsystem judgement.

Scope (Deviation 6 against pre-registration commit 823270e)
-----------------------------------------------------------
Coverage is extended from the 100 reliability items to every item the
annotator marked relevant at Stage 1. Rationale: six categories estimated
on ~59 unique relevant items is too thin to support RQ2's ranking claim,
the marginal cost is ~90 extra judgements per annotator, and the extension
was forced by sheet construction rather than chosen after seeing results.
A seventh value, `none`, is added because a comment can be CAV-relevant
without engaging any identifiable subsystem; forcing such comments into
one of six would manufacture signal. An optional `subsystem_secondary`
records multi-subsystem engagement without changing the primary variable,
so Krippendorff's alpha still runs on a single nominal variable as
pre-registered.

Design choices
--------------
* Each annotator labels only the items they themselves marked relevant.
  On the reliability items this yields an incomplete design where the
  three disagreed on relevance. Krippendorff's alpha handles that; it is
  one reason the statistic was chosen.
* The annotator's own Stage 2 class is shown as context. They are deciding
  which subsystem a comment engages, not re-deciding the class, and hiding
  their own prior label would be artificial.
* Original row order and row numbers are preserved, so the reliability
  items stay interleaved and the annotators remain blind to which items
  are double-rated.
* Encoding repair: sheet C was saved through a non-UTF-8 round trip and
  carries mojibake in a subset of comments. It is reversible exactly and
  is repaired here so the subsystem pass reads clean text.

Usage
-----
    python build_subsystem_retrofit.py

Reads   data/annotation/tranche1_annotator_{A,B,C}.{xlsx,csv}
        data/annotation/tranche1_master_key.csv          (verified if present)
Writes  data/annotation/subsystem/tranche1_subsystem_{A,B,C}.xlsx
        data/annotation/subsystem/tranche1_subsystem_effort.csv
        data/annotation/subsystem/subsystem_retrofit_manifest.txt
"""

from __future__ import annotations

import hashlib
import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

# --------------------------------------------------------------------------
# configuration
# --------------------------------------------------------------------------

ANNOTATORS = ["A", "B", "C"]

IN_DIR = Path("data/annotation")
OUT_DIR = Path("data/annotation/subsystem")
MASTER_KEY = IN_DIR / "tranche1_master_key.csv"

# Six pre-registered subsystems plus `none`. Check these tokens against the
# `subsystem` column of the calibration sheets before distributing; the
# script prints the calibration vocabulary if those files are found.
SUBSYSTEMS = [
    "sensing",
    "in_vehicle_networks",
    "v2x",
    "cybersecurity",
    "privacy",
    "data_governance",
    "none",
]
SECONDARY = [s for s in SUBSYSTEMS if s != "none"]

STAGE1_VOCAB = {"Y", "N"}
STAGE2_VOCAB = {"CONCERN", "ENDORSEMENT", "OTHER"}

EXPECTED_TOTAL_ITEMS = 500
EXPECTED_RELIABILITY = 100
EXPECTED_RATINGS = 700

MOJIBAKE_MARKERS = "‚Ä Â √ ï ì î"  # spaced for readability; split below
_MARK = [m for m in MOJIBAKE_MARKERS.split(" ") if m]


# --------------------------------------------------------------------------
# helpers
# --------------------------------------------------------------------------

def fail(msg: str) -> None:
    print(f"\n[FAIL] {msg}")
    sys.exit(1)


def membership_sha256(ids) -> str:
    """SHA-256 over sorted ids, newline-joined, no trailing newline."""
    payload = "\n".join(sorted(ids)).encode("utf-8")
    return hashlib.sha256(payload).hexdigest()


def load_sheet(letter: str) -> pd.DataFrame:
    xlsx = IN_DIR / f"tranche1_annotator_{letter}.xlsx"
    csv = IN_DIR / f"tranche1_annotator_{letter}.csv"
    if xlsx.exists():
        df = pd.read_excel(xlsx, sheet_name=0, dtype=str)
        src = xlsx
    elif csv.exists():
        df = pd.read_csv(csv, dtype=str, encoding="utf-8")
        src = csv
    else:
        fail(f"no completed sheet found for annotator {letter} "
             f"(looked for {xlsx} and {csv})")
    df = df.fillna("")
    for col in ("id", "subreddit", "stage1_relevant", "stage2_class", "skip"):
        if col in df.columns:
            df[col] = df[col].str.strip()
    print(f"   {letter}: {len(df):>4} rows from {src}")
    return df


def demojibake(text: str) -> str:
    """Reverse a UTF-8 -> mac_roman/cp1252 mis-decode. Returns text unchanged
    if no round trip both succeeds and removes the marker characters."""
    if not any(m in text for m in _MARK):
        return text
    for enc in ("mac_roman", "cp1252", "latin-1"):
        try:
            repaired = text.encode(enc).decode("utf-8")
        except (UnicodeEncodeError, UnicodeDecodeError):
            continue
        if not any(m in repaired for m in _MARK):
            return repaired
    return text


# --------------------------------------------------------------------------
# load and repair
# --------------------------------------------------------------------------

def main() -> None:
    print("=" * 74)
    print("tranche 1 subsystem retrofit")
    print("=" * 74)

    if not IN_DIR.exists():
        fail(f"{IN_DIR} does not exist - run this from the repository root")

    print("\n[load]")
    sheets = {x: load_sheet(x) for x in ANNOTATORS}

    required = ["row", "id", "subreddit", "comment",
                "stage1_relevant", "stage2_class", "skip", "notes"]
    for x, df in sheets.items():
        missing = [c for c in required if c not in df.columns]
        if missing:
            fail(f"annotator {x} sheet is missing column(s): {missing}")

    print("\n[encoding repair]")
    repaired_counts = {}
    for x, df in sheets.items():
        before = df["comment"].tolist()
        after = [demojibake(t) for t in before]
        n = sum(1 for b, a in zip(before, after) if b != a)
        still = sum(1 for a in after if any(m in a for m in _MARK))
        df["comment"] = after
        repaired_counts[x] = n
        note = f"   {x}: repaired {n:>3} comment(s)"
        if still:
            note += f"   WARNING: {still} row(s) still carry marker characters"
        print(note)

    # ----------------------------------------------------------------------
    # integrity checks
    # ----------------------------------------------------------------------

    print("\n[integrity]")
    ids = {x: set(df["id"]) for x, df in sheets.items()}

    for x, df in sheets.items():
        if len(df) != df["id"].nunique():
            fail(f"annotator {x} has duplicate ids within the sheet")

    union = set().union(*ids.values())
    shared = ids["A"] & ids["B"] & ids["C"]
    counts = pd.Series([i for s in ids.values() for i in s]).value_counts()
    n_single = int((counts == 1).sum())
    n_triple = int((counts == 3).sum())
    n_other = int(((counts != 1) & (counts != 3)).sum())
    total_ratings = sum(len(df) for df in sheets.values())

    def check(label, actual, expected):
        ok = actual == expected
        print(f"   {label:<34} actual {actual:>5}  expected {expected:>5}  "
              f"{'PASS' if ok else 'FAIL'}")
        return ok

    ok = True
    ok &= check("distinct items", len(union), EXPECTED_TOTAL_ITEMS)
    ok &= check("three-way reliability items", n_triple, EXPECTED_RELIABILITY)
    ok &= check("singly rated items", n_single,
                EXPECTED_TOTAL_ITEMS - EXPECTED_RELIABILITY)
    ok &= check("items rated twice (should be 0)", n_other, 0)
    ok &= check("total ratings", total_ratings, EXPECTED_RATINGS)
    if not ok:
        fail("reliability design does not match the pre-registered structure")

    # label vocabulary and completeness
    for x, df in sheets.items():
        bad1 = sorted(set(df["stage1_relevant"]) - STAGE1_VOCAB)
        if bad1:
            fail(f"annotator {x} has out-of-vocabulary Stage 1 value(s): {bad1}")
        bad2 = sorted(set(df["stage2_class"]) - STAGE2_VOCAB - {""})
        if bad2:
            fail(f"annotator {x} has out-of-vocabulary Stage 2 value(s): {bad2}")
        y_no_class = int(((df["stage1_relevant"] == "Y") &
                          (df["stage2_class"] == "")).sum())
        n_with_class = int(((df["stage1_relevant"] == "N") &
                            (df["stage2_class"] != "")).sum())
        if y_no_class or n_with_class:
            fail(f"annotator {x}: {y_no_class} relevant row(s) without a "
                 f"class, {n_with_class} irrelevant row(s) with a class")
    print("   label vocabulary and completeness          PASS")

    # shared items must now show identical text across the three sheets
    text = {x: dict(zip(df["id"], df["comment"])) for x, df in sheets.items()}
    sub = {x: dict(zip(df["id"], df["subreddit"])) for x, df in sheets.items()}
    txt_mismatch = [i for i in sorted(shared)
                    if not (text["A"][i] == text["B"][i] == text["C"][i])]
    sr_mismatch = [i for i in sorted(shared)
                   if not (sub["A"][i] == sub["B"][i] == sub["C"][i])]
    print(f"   shared items, text mismatches              "
          f"{len(txt_mismatch)}  {'PASS' if not txt_mismatch else 'FAIL'}")
    print(f"   shared items, subreddit mismatches         "
          f"{len(sr_mismatch)}  {'PASS' if not sr_mismatch else 'FAIL'}")
    if txt_mismatch:
        print("      first few:", txt_mismatch[:5])
        fail("shared reliability items do not carry identical text after "
             "encoding repair - do not proceed until this is explained")
    if sr_mismatch:
        fail("shared reliability items disagree on subreddit")

    skips = {x: int((df["skip"].str.strip() != "").sum())
             for x, df in sheets.items()}
    print(f"   skips recorded                             {skips}")

    # ----------------------------------------------------------------------
    # master key verification
    # ----------------------------------------------------------------------

    print("\n[master key]")
    if MASTER_KEY.exists():
        key = pd.read_csv(MASTER_KEY, dtype=str).fillna("")
        id_col = next((c for c in key.columns if c.lower() == "id"), None)
        if id_col is None:
            fail(f"{MASTER_KEY} has no 'id' column; columns are "
                 f"{list(key.columns)}")
        key_ids = set(key[id_col].str.strip())
        missing = key_ids - union
        extra = union - key_ids
        print(f"   key items                {len(key_ids)}")
        print(f"   in key, not in sheets    {len(missing)}")
        print(f"   in sheets, not in key    {len(extra)}")
        if missing or extra:
            if missing:
                print("      missing:", sorted(missing)[:10])
            if extra:
                print("      extra  :", sorted(extra)[:10])
            fail("sheet ids do not reconcile with the tranche 1 master key")
        print("   reconciliation           PASS")
        key_sha = membership_sha256(key_ids)
        print(f"   key membership sha256    {key_sha}")
    else:
        key_sha = None
        print(f"   {MASTER_KEY} not found - skipping reconciliation")
        print("   WARNING: emitted sheets are unverified against the draw")

    # calibration vocabulary cross-check, if the calibration sheets are here
    cal_vocab = set()
    for pattern in ("calibration_annotator_*.xlsx", "calibration_annotator_*.csv",
                    "calibration_*.xlsx", "calibration_*.csv"):
        for path in sorted(IN_DIR.glob(pattern)):
            try:
                cal = (pd.read_excel(path, dtype=str) if path.suffix == ".xlsx"
                       else pd.read_csv(path, dtype=str))
            except Exception:
                continue
            col = next((c for c in cal.columns if c.lower() == "subsystem"), None)
            if col is not None:
                cal_vocab |= {v.strip() for v in cal[col].fillna("")
                              if str(v).strip()}
    if cal_vocab:
        print(f"\n[calibration vocabulary] observed: {sorted(cal_vocab)}")
        unknown = sorted(cal_vocab - set(SUBSYSTEMS))
        if unknown:
            print(f"   WARNING: calibration used value(s) not in SUBSYSTEMS: "
                  f"{unknown}")
            print("   Reconcile the vocabulary before distributing the sheets.")
    else:
        print("\n[calibration vocabulary] calibration sheets not found - "
              "confirm SUBSYSTEMS matches the codebook by hand")

    # ----------------------------------------------------------------------
    # emit
    # ----------------------------------------------------------------------

    print("\n[emit]")
    OUT_DIR.mkdir(parents=True, exist_ok=True)

    header = ["row", "id", "subreddit", "comment", "stage2_class",
              "subsystem", "subsystem_secondary", "notes"]
    head_fill = PatternFill("solid", fgColor="DDDDDD")
    entry_fill = PatternFill("solid", fgColor="FFF6D5")
    widths = {"row": 6, "id": 12, "subreddit": 18, "comment": 95,
              "stage2_class": 14, "subsystem": 22,
              "subsystem_secondary": 22, "notes": 40}

    manifest_lines = []
    emitted = {}

    for x in ANNOTATORS:
        df = sheets[x]
        rel = df[df["stage1_relevant"] == "Y"].copy()
        rel = rel.sort_values("row", key=lambda s: s.astype(int))

        wb = Workbook()
        ws = wb.active
        ws.title = f"tranche1_subsystem_{x}"

        ws.append(header)
        for c in range(1, len(header) + 1):
            cell = ws.cell(row=1, column=c)
            cell.font = Font(bold=True)
            cell.fill = head_fill
            cell.alignment = Alignment(vertical="center")

        for _, r in rel.iterrows():
            ws.append([int(r["row"]), r["id"], r["subreddit"], r["comment"],
                       r["stage2_class"], "", "", ""])

        n = len(rel)
        last = n + 1

        for name, width in widths.items():
            col = get_column_letter(header.index(name) + 1)
            ws.column_dimensions[col].width = width
        ws.freeze_panes = "A2"

        for row in ws.iter_rows(min_row=2, max_row=last,
                                min_col=1, max_col=len(header)):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

        col_sub = get_column_letter(header.index("subsystem") + 1)
        col_sec = get_column_letter(header.index("subsystem_secondary") + 1)
        col_note = get_column_letter(header.index("notes") + 1)
        for col in (col_sub, col_sec, col_note):
            for rr in range(2, last + 1):
                ws[f"{col}{rr}"].fill = entry_fill

        dv_primary = DataValidation(
            type="list", formula1='"' + ",".join(SUBSYSTEMS) + '"',
            allow_blank=False, showDropDown=False)
        dv_primary.error = ("Choose one of: " + ", ".join(SUBSYSTEMS))
        dv_primary.errorTitle = "Not a valid subsystem"
        dv_primary.prompt = ("Which subsystem does this comment engage? "
                             "Choose 'none' if it is CAV-relevant but engages "
                             "no identifiable subsystem.")
        dv_primary.promptTitle = "Primary subsystem"
        ws.add_data_validation(dv_primary)
        dv_primary.add(f"{col_sub}2:{col_sub}{last}")

        dv_secondary = DataValidation(
            type="list", formula1='"' + ",".join(SECONDARY) + '"',
            allow_blank=True, showDropDown=False)
        dv_secondary.error = ("Leave blank, or choose one of: "
                              + ", ".join(SECONDARY))
        dv_secondary.errorTitle = "Not a valid subsystem"
        dv_secondary.prompt = ("Optional. Only if the comment clearly engages "
                               "a second subsystem as well. Leave blank "
                               "otherwise.")
        dv_secondary.promptTitle = "Secondary subsystem"
        ws.add_data_validation(dv_secondary)
        dv_secondary.add(f"{col_sec}2:{col_sec}{last}")

        out = OUT_DIR / f"tranche1_subsystem_{x}.xlsx"
        wb.save(out)
        emitted[x] = list(rel["id"])
        sha = membership_sha256(emitted[x])
        print(f"   {x}: {n:>4} items -> {out}")
        print(f"      membership sha256 {sha}")
        manifest_lines.append(f"annotator {x}: {n} items  sha256 {sha}")

    # effort template
    effort = OUT_DIR / "tranche1_subsystem_effort.csv"
    pd.DataFrame({
        "annotator": ANNOTATORS,
        "items": [len(emitted[x]) for x in ANNOTATORS],
        "start_datetime": ["", "", ""],
        "finish_datetime": ["", "", ""],
        "total_minutes": ["", "", ""],
        "notes": ["", "", ""],
    }).to_csv(effort, index=False)
    print(f"   effort template -> {effort}")

    # manifest
    all_rel = sorted(set().union(*[set(v) for v in emitted.values()]))
    manifest = OUT_DIR / "subsystem_retrofit_manifest.txt"
    with manifest.open("w", encoding="utf-8") as fh:
        fh.write("tranche 1 subsystem retrofit manifest\n")
        fh.write("=" * 60 + "\n\n")
        fh.write("Deviation 6 against pre-registration commit 823270e:\n")
        fh.write("  coverage extended from the 100 reliability items to every\n")
        fh.write("  item the annotator marked relevant; `none` added as a\n")
        fh.write("  seventh primary value; optional secondary subsystem added.\n\n")
        fh.write(f"vocabulary (primary):   {', '.join(SUBSYSTEMS)}\n")
        fh.write(f"vocabulary (secondary): {', '.join(SECONDARY)} or blank\n\n")
        for line in manifest_lines:
            fh.write(line + "\n")
        fh.write(f"\nunique relevant items across sheets: {len(all_rel)}\n")
        fh.write(f"union membership sha256: {membership_sha256(all_rel)}\n")
        fh.write(f"total subsystem judgements requested: "
                 f"{sum(len(v) for v in emitted.values())}\n")
        fh.write(f"\nreliability items carried into the retrofit: "
                 f"{sum(1 for i in emitted['A'] if i in shared)} (A), "
                 f"{sum(1 for i in emitted['B'] if i in shared)} (B), "
                 f"{sum(1 for i in emitted['C'] if i in shared)} (C)\n")
        unanimous_y = sorted(i for i in shared
                             if all(i in emitted[x] for x in ANNOTATORS))
        fh.write(f"reliability items relevant to all three (complete "
                 f"subsystem design): {len(unanimous_y)}\n")
        fh.write(f"\nencoding repairs applied: {repaired_counts}\n")
        if key_sha:
            fh.write(f"master key membership sha256: {key_sha}\n")
        else:
            fh.write("master key: NOT VERIFIED (file absent)\n")
    print(f"   manifest -> {manifest}")

    print("\n[summary]")
    print(f"   unique relevant items                {len(all_rel)}")
    print(f"   total subsystem judgements requested "
          f"{sum(len(v) for v in emitted.values())}")
    print(f"   reliability items relevant to all 3  "
          f"{sum(1 for i in shared if all(i in emitted[x] for x in ANNOTATORS))}")
    print("\n[done] sheets are ready to distribute.")


if __name__ == "__main__":
    main()
