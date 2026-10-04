"""Tranche 2, step 4: the draw, the researcher-only key and the annotator workbooks.

Implements R1 and Deviation 10.

  Eligible     the 12,166 frame items less the 500 tranche 1 items and the
               calibration items, as validated and frozen by step 1.
  Draw         for each predicted label (CONCERN, ENDORSEMENT, OTHER), a simple
               random sample without replacement of the allocated number from the
               eligible items with that predicted label. Items predicted not relevant
               are never drawn. No era stratification. No selection on confidence
               beyond the relevance threshold already in the predicted label.
  Reliability  100 of the 500, spread across predicted labels in proportion to
               the allocation (largest remainder, exact arithmetic, ties in the
               order CONCERN, ENDORSEMENT, OTHER), random within label.
  Assignment   the other 400 are shuffled within predicted label and dealt
               A, B, C in turn, so each annotator gets the same mix.
  Sheets       each annotator's 100 reliability items and own items in that
               annotator's own random order. No predicted label, score or
               reliability flag is shown. The text is the text that was scored.
               The two subsystem columns sit beside the Stage 2 class and are
               hidden until the annotator has finished Stage 1 and Stage 2.
  Random pool  a simple random sample of 500 eligible items, drawn without
               reference to the predicted label, for R1's confidence comparison.
               It is not annotated and may overlap tranche 2.

Every random step uses numpy.random.default_rng([24916660, 2, component]).

Within the predicted-relevant eligible items the draw is a stratified simple random
sample, so an item with predicted label p has inclusion probability a_p / N_p. Items
predicted not relevant have probability zero, so the tranche cannot estimate
prevalence in the frame and, under R1, is never used for it.

The draw is made once. Everything is written to a temporary folder and moved into
place only when every file has been written and read back, so a failure leaves
nothing behind. A completed draw is never overwritten.

SEALED OUTPUT. Like step 3, this script prints only checks and SHA-256 values.
"""
from __future__ import annotations

import csv
import shutil
import time

import numpy as np
import pandas as pd

import tranche2_common as c


def patiently(action, attempts: int = 20, wait: float = 0.5):
    """Run a file operation, waiting out a short lock. On Windows a virus scanner or a
    sync client (the repository may sit in OneDrive) can hold a new file for a moment."""
    for k in range(attempts):
        try:
            return action()
        except OSError:                             # PermissionError (file in use), "directory not empty", ...
            if k == attempts - 1:
                raise
            time.sleep(wait)


def largest_remainder(total: int, weights: dict[str, int]) -> dict[str, int]:
    """Proportional split in exact integer arithmetic; ties go to the earlier key."""
    s = sum(weights.values())
    out = {k: (total * v) // s for k, v in weights.items()}
    rem = {k: (total * v) % s for k, v in weights.items()}
    left = total - sum(out.values())
    keys = list(weights)
    for k in sorted(keys, key=lambda k: (-rem[k], keys.index(k)))[:left]:
        out[k] += 1
    return out


def write_workbook(path, annotator: str, rows: list[dict]) -> None:
    """rows: dicts with id, subreddit, comment, in sheet order."""
    import math

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    header = ["row", "id", "subreddit", "comment", "stage1_relevant", "stage2_class"]
    widths = [6, 11, 17, 78, 16, 15]
    second_pass: list[str] = []
    if c.INCLUDE_SUBSYSTEM_COLUMNS:                 # second-pass columns: beside the class, hidden until pass 1 is done
        second_pass = ["subsystem", "subsystem_secondary"]
        header += second_pass
        widths += [21, 21]
    header += ["skip", "notes"]
    widths += [6, 40]
    per_line = max(40, int(widths[3] * 1.05))       # characters of comment text per wrapped line, roughly

    wb = Workbook()
    ws = wb.active
    ws.title = f"tranche2_annotator_{annotator}"
    ws.append(header)
    for j, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
        cell = ws.cell(row=1, column=j)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="BDD7EE" if header[j - 1] in second_pass else "DDDDDD")
        cell.alignment = Alignment(vertical="top")
    top = Alignment(vertical="top")
    wrap = Alignment(vertical="top", wrap_text=True)
    for n, r in enumerate(rows, start=1):
        text = r["comment"]
        ws.cell(row=n + 1, column=1, value=n).alignment = Alignment(vertical="top", horizontal="center")
        for col, value in ((2, r["id"]), (3, r["subreddit"]), (4, text)):
            cell = ws.cell(row=n + 1, column=col)
            cell.value = value
            cell.data_type = "s"                    # stored as text: never read as a formula ("=...") or a number ("1e5")
            cell.alignment = wrap if col == 4 else top
        for col in range(5, len(header) + 1):       # cell formats stay General, as in the tranche 1 sheets
            cell = ws.cell(row=n + 1, column=col)
            cell.alignment = wrap if header[col - 1] == "notes" else top
        lines = max(1, math.ceil(len(text) / per_line) + text.count("\n"))
        ws.row_dimensions[n + 1].height = min(409, 15 * lines + 3)
    ws.freeze_panes = "A2"                          # the header row stays in view; no column is pinned (as in tranche 1)
    if second_pass:                                 # collapsed group: one click on the + above the columns reveals them
        first = get_column_letter(header.index(second_pass[0]) + 1)
        final = get_column_letter(header.index(second_pass[-1]) + 1)
        ws.column_dimensions.group(first, final, hidden=True)

    last = len(rows) + 1

    def dropdown(column_name: str, options: list[str]) -> None:
        col = get_column_letter(header.index(column_name) + 1)
        dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True,
                            showErrorMessage=True, errorTitle="Not a valid entry",
                            error="Choose one of: " + ", ".join(options))
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{last}")

    dropdown("stage1_relevant", ["Y", "N"])
    dropdown("stage2_class", c.CLASSES)
    dropdown("skip", ["Y"])
    if c.INCLUDE_SUBSYSTEM_COLUMNS:
        dropdown("subsystem", c.SUBSYSTEMS + ["none"])
        dropdown("subsystem_secondary", c.SUBSYSTEMS)
    wb.save(path)


def quantiles(x: pd.Series) -> dict:
    q = x.astype(float).quantile([0.10, 0.25, 0.50, 0.75, 0.90])
    return {f"p{int(round(100 * k))}": round(float(v), 4) for k, v in q.items()}


def main() -> None:
    c.setup_console()
    print("Tranche 2, step 4: draw, key and workbooks\n")
    final = {"key": c.MASTER_KEY, "manifest": c.DRAW_MANIFEST, "pool": c.RANDOM_POOL_IDS, "figures": c.SEALED_DRAW_TEXT}
    final.update({a: c.workbook_path(a) for a in c.ANNOTATORS})
    existing = [p for p in final.values() if p.exists()]
    if existing:
        c.stop("tranche 2 output already exists and a frozen draw is never overwritten:\n  "
               + "\n  ".join(str(p) for p in existing)
               + "\n  If a redraw is genuinely needed, record the reason as a deviation first.")
    if not c.ALLOCATION_JSON.exists():
        c.stop("run step 3 first: tranche2_allocation.json is missing.")

    step1 = c.step1_manifest()
    frame_ids = c.load_frame()
    corpus, _ = c.load_corpus(frame_ids)
    corpus = corpus.set_index("id")
    key1, _, _ = c.load_tranche1(verbose=False)
    calibration = c.load_resolved_calibration(step1)
    alloc = c.read_json(c.ALLOCATION_JSON)
    if alloc.get("gate", {}).get("result") != "PASS":
        c.stop("the gate in step 3 did not pass. Nothing is drawn under Deviation 10.")

    scores_sha = c.file_sha256(c.PREFILTER_SCORES)
    if scores_sha != alloc["inputs"]["scores_sha256"]:
        c.stop("prefilter_scores.csv is not the file the allocation was computed from.")
    if alloc["inputs"].get("calibration_ids_sha256") != c.ids_sha256(calibration):
        c.stop("the calibration ids are not those the allocation was computed with.")
    scores = pd.read_csv(c.PREFILTER_SCORES, dtype={"id": str}, keep_default_na=False, encoding="utf-8").set_index("id")
    eligible = c.eligible_ids(frame_ids, key1.index, calibration, step1)
    if alloc["inputs"].get("eligible_ids_sha256") != step1["eligible_ids_sha256"]:
        c.stop("the eligible items are not those the allocation was computed with.")
    pools = {k: sorted(i for i in eligible if scores.at[i, "pred_label"] == k) for k in c.LABELS}
    if {k: len(v) for k, v in pools.items()} != alloc["eligible_by_predicted_label"]:
        c.stop("the eligible pools differ from those recorded in tranche2_allocation.json.")
    allocation = {k: int(alloc["allocation"][k]) for k in c.CLASSES}
    if sum(allocation.values()) != c.TRANCHE2_N:
        c.stop("the allocation does not sum to 500.")
    print(f"allocation read from {c.ALLOCATION_JSON.name} (not displayed); gate PASS")

    # the text shown to annotators is the text that was scored
    if c.file_sha256(c.PREFILTER_INPUT) != step1["file_sha256"]:
        c.stop("prefilter_input.csv has changed since step 1 wrote it.")
    scored = pd.read_csv(c.PREFILTER_INPUT, dtype=str, keep_default_na=False, encoding="utf-8").set_index("id")["text"]
    if not (scored.reindex(corpus.index) == corpus["text"]).all():
        c.stop("the corpus text no longer matches the text that was scored. The corpus file has changed since step 1.")

    # ---- draw ----------------------------------------------------------------------
    drawn: dict[str, list[str]] = {}
    for k in c.CLASSES:
        n = allocation[k]
        if n > len(pools[k]):
            c.stop("an allocation exceeds its pool")
        pick = c.rng(f"draw_{k}").choice(len(pools[k]), size=n, replace=False) if n else []
        drawn[k] = [pools[k][j] for j in sorted(int(j) for j in pick)]
    ids = sorted(i for k in c.CLASSES for i in drawn[k])
    pred = {i: k for k in c.CLASSES for i in drawn[k]}

    # ---- reliability subset --------------------------------------------------------
    quota = largest_remainder(c.RELIABILITY_N, {k: len(drawn[k]) for k in c.CLASSES if drawn[k]})
    g = c.rng("reliability")
    reliability: set[str] = set()
    for k in c.CLASSES:
        if not drawn[k]:
            continue
        pick = g.choice(len(drawn[k]), size=quota[k], replace=False)
        reliability |= {drawn[k][int(j)] for j in pick}

    # ---- assignment of the remaining items ------------------------------------------
    g = c.rng("assignment")
    sequence: list[str] = []
    for k in c.CLASSES:
        singles = [i for i in drawn[k] if i not in reliability]
        sequence += [singles[int(j)] for j in g.permutation(len(singles))]
    assigned = {i: c.ANNOTATORS[pos % len(c.ANNOTATORS)] for pos, i in enumerate(sequence)}

    # ---- sheets ---------------------------------------------------------------------
    sheets: dict[str, list[str]] = {}
    for a in c.ANNOTATORS:
        items = sorted(reliability | {i for i, who in assigned.items() if who == a})
        order = c.rng(f"order_{a}").permutation(len(items))
        sheets[a] = [items[int(j)] for j in order]

    # ---- random pool for the bias analysis ------------------------------------------
    pick = c.rng("random_pool").choice(len(eligible), size=c.RANDOM_POOL_N, replace=False)
    random_pool = sorted(eligible[int(j)] for j in pick)

    # ---- self-checks before anything is written -------------------------------------
    t1, cal = set(key1.index), set(calibration)
    too_long = [i for i in ids if c.excel_len(scored[i]) > c.EXCEL_CELL_LIMIT]
    unclean = [i for i in ids if c.shown_text(scored[i]) != scored[i]]
    checks = {
        "500 distinct ids": len(ids) == c.TRANCHE2_N == len(set(ids)),
        "all eligible (in frame, not tranche 1, not calibration)": set(ids) <= set(eligible) and not (set(ids) & (t1 | cal)),
        "every item is predicted relevant": all(scores.at[i, "pred_label"] in c.CLASSES for i in ids),
        "counts by predicted label equal the allocation": all(len(drawn[k]) == allocation[k] for k in c.CLASSES),
        "100 reliability items": len(reliability) == c.RELIABILITY_N,
        "400 single items, each assigned once": len(assigned) == c.TRANCHE2_N - c.RELIABILITY_N
                                                 and not (set(assigned) & reliability),
        "reliability items in every sheet": all(reliability <= set(sheets[a]) for a in c.ANNOTATORS),
        "700 ratings in total": sum(len(v) for v in sheets.values()) == c.TRANCHE2_N + 2 * c.RELIABILITY_N,
        "random pool: 500 distinct eligible ids": len(set(random_pool)) == c.RANDOM_POOL_N and set(random_pool) <= set(eligible),
        "every comment fits in a workbook cell and is clean": not too_long and not unclean,
    }
    for name, ok in checks.items():
        print(f"  [{'ok' if ok else 'FAIL'}] {name}")
    if not all(checks.values()):
        c.stop("a self-check failed; nothing was written.")

    # ---- write to a temporary folder, verify, then move into place --------------------
    tmp = c.OUT_DIR / "_draw_in_progress"
    if tmp.exists():
        patiently(lambda: shutil.rmtree(tmp))
    patiently(lambda: tmp.mkdir(parents=True))
    try:
        key = pd.DataFrame({
            "id": ids,
            "era": [corpus.at[i, "era"] for i in ids],
            "year": [int(corpus.at[i, "year"]) for i in ids],
            "tier": [corpus.at[i, "tier"] for i in ids],
            "subreddit": [corpus.at[i, "subreddit"] for i in ids],
            "is_reliability": [i in reliability for i in ids],
            "assigned_to": [assigned.get(i, "") for i in ids],
            "tranche": c.TRANCHE,
            "pred_label": [pred[i] for i in ids],
            "p_relevant": [float(scores.at[i, "p_relevant"]) for i in ids],
            "confidence": [float(scores.at[i, "confidence"]) for i in ids],
        })
        key.to_csv(tmp / final["key"].name, index=False, encoding="utf-8-sig", lineterminator="\n",
                   float_format="%.6f", quoting=csv.QUOTE_MINIMAL)

        for a in c.ANNOTATORS:
            rows = [{"id": i, "subreddit": corpus.at[i, "subreddit"], "comment": scored[i]} for i in sheets[a]]
            path = tmp / final[a].name
            write_workbook(path, a, rows)
            back = pd.read_excel(path, dtype=object, keep_default_na=False)       # read back what was written
            if back["id"].astype(str).tolist() != sheets[a] or back["row"].tolist() != list(range(1, len(rows) + 1)):
                raise RuntimeError(f"workbook {a} did not read back as written")
            same_text = sum(c.norm_ws(x) == c.norm_ws(r["comment"]) for x, r in zip(back["comment"], rows))
            if same_text != len(rows):
                raise RuntimeError(f"workbook {a}: {len(rows) - same_text} comments did not read back as written")
        with open(tmp / final["pool"].name, "w", encoding="utf-8", newline="\n") as f:
            f.write("\n".join(random_pool) + "\n")

        by_label = lambda members: {k: int(sum(pred[i] == k for i in members)) for k in c.CLASSES}
        eras = key["era"].value_counts().sort_index()
        manifest = {
            "created": c.now_iso(),
            "specification": "R1 and Deviation 10, docs/analysis_preregistration.md",
            "seed": c.SEED,
            "rng": "numpy.random.default_rng([seed, tranche, component])",
            "rng_components": c.RNG_COMPONENT,
            "allocation": allocation,
            "tranche2_ids_sha256": c.ids_sha256(ids),
            "reliability_ids_sha256": c.ids_sha256(reliability),
            "random_pool_ids_sha256": c.ids_sha256(random_pool),
            "random_pool_overlap_with_tranche2": len(set(random_pool) & set(ids)),
            "annotators": {a: {"items": len(sheets[a]), "own_items": len(sheets[a]) - c.RELIABILITY_N,
                               "ids_sha256": c.ids_sha256(sheets[a]),
                               "own_by_predicted_label": by_label([i for i in sheets[a] if i not in reliability]),
                               "workbook": final[a].name,
                               "workbook_sha256": c.file_sha256(tmp / final[a].name)} for a in c.ANNOTATORS},
            "reliability_by_predicted_label": by_label(reliability),
            "tranche2_by_era": {k: int(v) for k, v in eras.items()},
            "tranche2_by_subreddit": {k: int(v) for k, v in key["subreddit"].value_counts().sort_index().items()},
            "confidence_comparison": {
                name: {"p_relevant": quantiles(scores.loc[members, "p_relevant"]),
                       "class_confidence": quantiles(scores.loc[members, "confidence"])}
                for name, members in (("tranche2", ids), ("random_pool", random_pool), ("all_eligible", eligible))},
            "subsystem_columns_in_workbooks": bool(c.INCLUDE_SUBSYSTEM_COLUMNS),
            "master_key": final["key"].name,
            "master_key_sha256": c.file_sha256(tmp / final["key"].name),
            "inputs": {"allocation_json_sha256": c.file_sha256(c.ALLOCATION_JSON), "scores_sha256": scores_sha,
                       "prefilter_input_sha256": step1["file_sha256"],
                       "frame_sha256": c.FRAME_SHA256, "tranche1_ids_sha256": c.TRANCHE1_SHA256,
                       "calibration_ids_sha256": c.ids_sha256(calibration),
                       "eligible_ids_sha256": step1["eligible_ids_sha256"], "eligible_items": len(eligible)},
            "versions": c.versions(),
        }
        c.write_json(tmp / final["manifest"].name, manifest)
        m = manifest
        figures = [
            "### Tranche 2 draw: composition",
            "",
            "Unsealed under the Blinding clause of Deviation 10. Produced by `tranche2_04_draw.py`;",
            "the membership hashes were recorded at the draw.",
            "",
            "- By predicted label: " + ", ".join(f"{k} {allocation[k]}" for k in c.CLASSES) + ".",
            "- By era: " + ", ".join(f"{k} {v}" for k, v in m["tranche2_by_era"].items()) + ".",
            "- Reliability subset by predicted label: "
            + ", ".join(f"{k} {v}" for k, v in m["reliability_by_predicted_label"].items()) + ".",
            "- Own items by predicted label: " + "; ".join(
                f"{a} " + ", ".join(f"{k} {v}" for k, v in m["annotators"][a]["own_by_predicted_label"].items())
                for a in c.ANNOTATORS) + ".",
            f"- Random pool: {m['random_pool_overlap_with_tranche2']} of its 500 items are also in tranche 2.",
        ]
        with open(tmp / final["figures"].name, "w", encoding="utf-8", newline="\n") as f:
            f.write("\n".join(figures) + "\n")
        manifest_sha = c.file_sha256(tmp / final["manifest"].name)

        for p in final.values():                    # everything exists and was verified: move into place
            patiently(lambda p=p: (tmp / p.name).replace(p))
    except BaseException:
        shutil.rmtree(tmp, ignore_errors=True)
        left = [tmp] if tmp.exists() else []
        for p in final.values():                    # a failure part-way through the move leaves nothing behind
            try:
                patiently(lambda p=p: p.unlink(missing_ok=True), attempts=6)
            except OSError:
                left.append(p)
        if left:
            print("\nSTOPPED: writing the draw failed, and these could not be removed:")
            for p in left:
                print(f"  {p}")
            print("Delete them by hand without opening them, then run this step again. The draw is deterministic,")
            print("so the second run gives the same items.")
        else:
            print("\nSTOPPED: writing the draw failed; nothing was left behind. The draw is deterministic, so running")
            print("this step again after fixing the cause gives the same items.")
        raise
    shutil.rmtree(tmp, ignore_errors=True)

    print("\nwrote, in " + str(c.OUT_DIR) + ":")
    for a in c.ANNOTATORS:
        print(f"  {final[a].name}: {len(sheets[a])} items = {c.RELIABILITY_N} reliability + "
              f"{len(sheets[a]) - c.RELIABILITY_N} own")
    print(f"  {final['key'].name} (RESEARCHER ONLY)\n  {final['manifest'].name}\n"
          f"  {final['figures'].name}\n  {final['pool'].name}")
    print("\n" + c.SEALED_NOTICE)

    c.show_register_block(c.PASTE_STEP4, [
        f"- Tranche 2 drawn {m['created']} by `tranche2_04_draw.py`: 500 items, membership SHA-256",
        f"  `{m['tranche2_ids_sha256']}`.",
        f"- Reliability subset: 100 items, SHA-256 `{m['reliability_ids_sha256']}`.",
        "- Sheets, by membership SHA-256:",
        *[f"  {a}, {m['annotators'][a]['items']} items, `{m['annotators'][a]['ids_sha256']}`"
          + (";" if a != c.ANNOTATORS[-1] else ".") for a in c.ANNOTATORS],
        f"- Random pool for the confidence comparison: 500 items, SHA-256 `{m['random_pool_ids_sha256']}`.",
        f"- `tranche2_master_key.csv` SHA-256 `{m['master_key_sha256']}`;",
        f"  `tranche2_manifest.json` SHA-256 `{manifest_sha}`.",
        "  Neither file is committed to the repository; both are sealed under Deviation 10.",
        f"- Seed {c.SEED}; NumPy {np.__version__}; pandas {pd.__version__}.",
    ])
    print("\nSend each annotator only their own workbook, without opening B's or C's.")


if __name__ == "__main__":
    main()
